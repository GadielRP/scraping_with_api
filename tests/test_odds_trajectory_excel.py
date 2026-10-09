from dataclasses import replace
from datetime import datetime, timedelta, timezone
from decimal import Decimal
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import cm_to_EMU
import pytest

import modules.jobs.pre_start_check_job.debug_exports.odds_trajectory_excel as excel_exporter
from modules.jobs.pre_start_check_job.debug_exports.odds_trajectory_excel import (
    FINAL_COLUMN_WIDTHS,
    TRAJECTORY_CHART_ANCHOR_COLUMN,
    TRAJECTORY_CHART_LINE_WIDTH_EMU,
    TRAJECTORY_CHART_RESERVED_ROWS,
    TRAJECTORY_CHART_SIZE_CM,
    TRAJECTORY_COLUMN_WIDTHS,
    TRAJECTORY_SEPARATOR_END_COLUMN,
    TRAJECTORY_TEXT_MERGE_END_COLUMNS,
    ROW_HEIGHTS_PT,
    export_odds_trajectory_context_xlsx,
)
from modules.pillars.context import EventIdentity
from modules.pillars.odds_trajectory_context import (
    OddsTrajectoryContext, MarketLineOddsTrajectory, BookieOddsTrajectory,
    ChoiceOddsTrajectory, OddsPointMeta, OddsSnapshotPoint,
)

UTC = timezone.utc
EARLY = datetime(2026, 10, 8, 12, 0, tzinfo=UTC)
CURRENT = EARLY + timedelta(hours=2)


def choice(name, initial="1.90", current="1.85"):
    initial_value = Decimal(initial) if initial is not None else None
    value = Decimal(current)
    snapshots = [
        OddsSnapshotPoint(1, 10, Decimal("1.90"), EARLY + timedelta(minutes=10), EARLY, Decimal(120)),
        OddsSnapshotPoint(2, 10, Decimal("1.90"), EARLY + timedelta(minutes=20), EARLY + timedelta(minutes=1), Decimal(119)),
        OddsSnapshotPoint(3, 10, value, CURRENT + timedelta(minutes=10), CURRENT, Decimal(5)),
        # Later raw snapshot must not replace the selected target's source timestamp.
        OddsSnapshotPoint(4, 10, Decimal("1.75"), CURRENT + timedelta(minutes=20), CURRENT + timedelta(minutes=15), Decimal(0)),
    ]
    return ChoiceOddsTrajectory(
        name, 1, initial_value, 10,
        odds_values={120: Decimal("2.50"), 5: value},
        meta_by_minute={5: OddsPointMeta(3, CURRENT + timedelta(minutes=10), 5, 5, Decimal(0))},
        snapshots=snapshots,
    )


def context(specs, targets=(120, 5)):
    markets = {}
    for index, (group, period, line, name, side, choices) in enumerate(specs):
        bookie = BookieOddsTrajectory(index, name, exchange_side=side, choices={c.choice_name: c for c in choices})
        market = MarketLineOddsTrajectory(index, f"{group} {period}", group, period, line, {str(index): bookie})
        lines = markets.setdefault(group, {}).setdefault(period, {}).setdefault(market.market_name, {})
        existing = lines.get(line or "default")
        if existing is not None:
            existing.bookies[str(index)] = bookie
        else:
            lines[line or "default"] = market
    return OddsTrajectoryContext(True, 4004, list(targets), list(targets), [], markets)


def export(tmp_path, ctx):
    path = tmp_path / "trajectory.xlsx"
    event = EventIdentity(
        4004, "Home vs Away", datetime(2026, 8, 31, 18, 0, tzinfo=UTC),
        5, "Football",
    )
    export_odds_trajectory_context_xlsx(ctx, path, event_context=event)
    return load_workbook(path)


def test_market_sheets_pivot_choices_and_have_excel_features(tmp_path):
    ctx = context([
        ("1X2", "Full Time", None, "SofaScore", None, [choice("1"), choice("x"), choice("2")]),
        ("Asian Handicap", "Full Time", "0", "Pinnacle Sports", None, [choice("1"), choice("2")]),
        ("Over/Under", "1st Half", "2.5", "bet365", None, [choice("over"), choice("under")]),
    ])
    workbook = export(tmp_path, ctx)
    assert workbook.sheetnames == ["Forma actual 1X2", "Trayectoria 1X2", "Forma actual Asian Handicap", "Trayectoria Asian Handicap", "Forma actual Over-Under", "Trayectoria Over-Under"]
    for sheet, names in (("Forma actual 1X2", ("1", "X", "2")), ("Forma actual Asian Handicap", ("1", "2")), ("Forma actual Over-Under", ("Over", "Under"))):
        ws = workbook[sheet]
        assert ws.max_row == 4
        assert [cell.value for cell in ws[3]][:4] == ["Periodo", "Línea", "Bookmaker", "Lado Exchange"]
        assert [cell.value for cell in ws[3]][4:] == [f"{name} {label}" for name in names for label in ("Inicial", "Fecha y hora del cambio (Inicial)", "Actual", "Fecha y hora del cambio (Actual)")]
        assert ws.freeze_panes == "E4"
        assert ws.auto_filter.ref.startswith("A3:")
        assert ws["E4"].number_format == "0.###"
        assert ws["F4"].number_format == "yyyy-mm-dd hh:mm:ss"
        assert ws.column_dimensions["B"].width == FINAL_COLUMN_WIDTHS["Línea"]
        assert ws.column_dimensions["D"].width == FINAL_COLUMN_WIDTHS["Lado Exchange"]
        assert ws.column_dimensions["E"].width == FINAL_COLUMN_WIDTHS["choice_odds"]
        assert ws.column_dimensions["F"].width == FINAL_COLUMN_WIDTHS["change_timestamp"]
        assert "Inicio: 2026-08-31 18:00:00 UTC" in ws["A1"].value
        assert ws["A1"].value.startswith("Forma actual —")
        assert "source_collected_at (UTC)" in ws["A2"].value
        assert "Fecha y hora del cambio" in " ".join(cell.value for cell in ws[3])
    assert workbook["Forma actual Asian Handicap"]["B4"].value == 0
    workbook.close()


def test_initial_and_current_use_exact_source_timestamps(tmp_path):
    workbook = export(tmp_path, context([("1X2", "Full Time", None, "SofaScore", None, [choice("1")])]))
    ws = workbook["Forma actual 1X2"]
    assert ws["E4"].value == 1.9  # initial_odds, not T120's 2.5
    assert ws["F4"].value == EARLY.replace(tzinfo=None)
    assert ws["G4"].value == 1.85
    assert ws["H4"].value == CURRENT.replace(tzinfo=None)
    workbook.close()


def test_current_is_most_recent_present_target_even_when_it_is_not_five(tmp_path):
    c = choice("1")
    c.odds_values[30] = Decimal("2.10")
    c.meta_by_minute[30] = replace(c.meta_by_minute[5], target_minute=30, snapshot_id=2)
    c.odds_values[1] = Decimal("1.75")
    c.meta_by_minute[1] = replace(c.meta_by_minute[5], target_minute=1, snapshot_id=4)
    for targets, expected, timestamp in (((120, 30), 2.1, EARLY + timedelta(minutes=1)), ((5, 120, 1), 1.75, CURRENT + timedelta(minutes=15))):
        workbook = export(tmp_path, context([("1X2", "Full Time", None, "SofaScore", None, [c])], targets))
        assert workbook["Forma actual 1X2"]["G4"].value == expected
        assert workbook["Forma actual 1X2"]["H4"].value == timestamp.replace(tzinfo=None)
        assert f"T{min(targets)}" in workbook["Forma actual 1X2"]["A2"].value
        workbook.close()


def test_missing_prices_and_missing_source_dates_stay_blank(tmp_path):
    c = choice("1", initial=None)
    c.odds_values.pop(5)
    workbook = export(tmp_path, context([("1X2", "Full Time", None, "SofaScore", None, [c])]))
    assert [workbook["Forma actual 1X2"].cell(4, col).value for col in range(5, 9)] == [None] * 4
    workbook.close()
    c = choice("1", initial="1.90000000000000000001")
    c.snapshots[2] = replace(c.snapshots[2], source_collected_at=None)
    workbook = export(tmp_path, context([("1X2", "Full Time", None, "SofaScore", None, [c])]))
    assert workbook["Forma actual 1X2"]["F4"].value is None  # Decimal inequality; no float matching
    assert workbook["Forma actual 1X2"]["H4"].value is None  # Never substitute collected_at
    workbook.close()


def test_bookmakers_periods_lines_and_exchange_sides_have_deterministic_order(tmp_path):
    specs = [("Asian Handicap", "Full Time", "0", name, side, [choice("1"), choice("2")]) for name, side in (
        ("Zulu", None), ("Betfair Exchange", "lay"), ("bet365", None),
        ("Pinnacle Sports", None), ("SofaScore", None), ("Alpha", None), ("Betfair Exchange", "back"),
    )]
    specs.extend(("Asian Handicap", period, line, "SofaScore", None, [choice("1"), choice("2")]) for period, line in (("1st Half", "10"), ("1st Half", "2")))
    workbook = export(tmp_path, context(specs))
    rows = list(workbook["Forma actual Asian Handicap"].iter_rows(min_row=4, values_only=True))
    assert [row[2] for row in rows] == ["SofaScore"] * 3 + ["Pinnacle Sports", "bet365", "Betfair Exchange", "Betfair Exchange", "Alpha", "Zulu"]
    assert [(row[0], row[1]) for row in rows[:3]] == [("1st Half", 2), ("1st Half", 10), ("Full Time", 0)]
    assert [row[3] for row in rows[5:7]] == ["Back", "Lay"]
    section_titles = [
        cell.value for cell in workbook["Trayectoria Asian Handicap"]["A"]
        if isinstance(cell.value, str) and "Choice " in cell.value
    ]
    assert {name.split(" | ")[2] for name in section_titles} == {
        "SofaScore", "Pinnacle Sports", "bet365", "Betfair Exchange", "Alpha", "Zulu",
    }
    assert any(" | Back | " in name for name in section_titles)
    assert any(" | Lay | " in name for name in section_titles)

    trajectory_sheet = workbook["Trayectoria Asian Handicap"]
    bookmaker_headers = [
        cell.value.removeprefix("Bookmaker: ")
        for cell in trajectory_sheet["A"]
        if isinstance(cell.value, str) and cell.value.startswith("Bookmaker: ")
    ]
    assert bookmaker_headers == [
        "SofaScore", "Pinnacle Sports", "bet365", "Betfair Exchange", "Alpha", "Zulu",
    ]
    current_bookmaker = None
    bookmaker_header_count = 0
    for row_number, cell in enumerate(trajectory_sheet["A"], start=1):
        if isinstance(cell.value, str) and cell.value.startswith("Bookmaker: "):
            bookmaker_header_count += 1
            current_bookmaker = cell.value.removeprefix("Bookmaker: ")
            if bookmaker_header_count > 1:
                separator_row = row_number - 1
                assert trajectory_sheet.cell(separator_row, 1).value is None
                assert trajectory_sheet.row_dimensions[separator_row].height == ROW_HEIGHTS_PT[
                    "bookmaker_separator"
                ]
                assert all(
                    trajectory_sheet.cell(separator_row, column).fill.fgColor.rgb.endswith("FCE4D6")
                    for column in range(1, TRAJECTORY_SEPARATOR_END_COLUMN + 1)
                )
        elif isinstance(cell.value, str) and "Choice " in cell.value:
            assert cell.value.split(" | ")[2] == current_bookmaker
    workbook.close()


def test_unknown_market_sheet_names_are_safe_and_unique(tmp_path):
    group = "Future/Market:*?[With]VeryLongNameMoreThan31"
    workbook = export(tmp_path, context([
        (group, "New Period", None, "New Book", None, [choice("Yes")]),
        (group.replace("/", "-"), "New Period", None, "New Book", None, [choice("No")]),
    ]))
    assert len(workbook.sheetnames) == 4
    assert len(set(workbook.sheetnames)) == 4
    assert all(len(name) <= 31 and not any(char in name for char in "[]:*?/\\") for name in workbook.sheetnames)
    assert all(ws.max_row >= 4 for ws in workbook)
    workbook.close()


def test_empty_context_writes_a_readable_workbook(tmp_path):
    workbook = export(tmp_path, context([], targets=()))
    assert workbook.sheetnames == ["Forma actual Sin mercados", "Trayectoria Sin mercados"]
    assert "sin target presente" in workbook["Forma actual Sin mercados"]["A2"].value
    workbook.close()



def test_significant_trajectory_sheet_includes_initial_and_all_changes_by_minute(tmp_path):
    one, draw, two = choice("1"), choice("x"), choice("2")
    one.snapshots[1] = replace(one.snapshots[1], odds_value=Decimal("2.05"))
    one.snapshots.append(replace(one.snapshots[1], snapshot_id=5, odds_value=Decimal("2.15")))
    draw.snapshots[1] = replace(draw.snapshots[1], odds_value=Decimal("3.20"))
    two.snapshots[1] = replace(two.snapshots[1], odds_value=Decimal("4.10"))
    workbook = export(tmp_path, context([
        ("1X2", "Full Time", None, "SofaScore", None, [one, draw, two]),
    ]))
    ws = workbook["Trayectoria 1X2"]
    assert ws["A1"].value.startswith("Trayectoria significativa — 1X2")
    assert ws.row_dimensions[1].height == ROW_HEIGHTS_PT["trajectory_title"]
    assert ws.row_dimensions[2].height == ROW_HEIGHTS_PT["trajectory_note"]
    assert ws.column_dimensions["A"].width == TRAJECTORY_COLUMN_WIDTHS["Cambio"]
    assert ws.column_dimensions["B"].width == TRAJECTORY_COLUMN_WIDTHS["Minutos antes del inicio"]
    assert ws.column_dimensions["C"].width == TRAJECTORY_COLUMN_WIDTHS["Cuota"]
    title_end = get_column_letter(TRAJECTORY_TEXT_MERGE_END_COLUMNS["title"])
    note_end = get_column_letter(TRAJECTORY_TEXT_MERGE_END_COLUMNS["note"])
    assert {str(cell_range) for cell_range in ws.merged_cells.ranges} >= {
        f"A1:{title_end}1", f"A2:{note_end}2",
    }
    assert "Cada choice tiene su propia secuencia" in ws["A2"].value
    sections = {}
    for index, cell in enumerate(ws["A"], start=1):
        if isinstance(cell.value, str) and "Choice " in cell.value:
            choice_name = cell.value.rsplit("Choice ", 1)[1]
            assert [ws.cell(index + 1, column).value for column in (1, 2, 3)] == [
                "Cambio", "Minutos antes del inicio", "Cuota",
            ]
            points = []
            row = index + 2
            while ws.cell(row, 3).value is not None:
                points.append((ws.cell(row, 1).value, ws.cell(row, 2).value, ws.cell(row, 3).value))
                row += 1
            sections[choice_name] = points

    assert list(sections) == ["1", "X", "2"]
    assert sections["1"] == [
        ("Inicial", Decimal(120), 1.9), (1, Decimal(119), 2.05),
        (2, Decimal(119), 2.15), (3, Decimal(5), 1.85), (4, Decimal(0), 1.75),
    ]
    assert sections["X"] == [
        ("Inicial", Decimal(120), 1.9), (1, Decimal(119), 3.2),
        (2, Decimal(5), 1.85), (3, Decimal(0), 1.75),
    ]
    assert sections["2"] == [
        ("Inicial", Decimal(120), 1.9), (1, Decimal(119), 4.1),
        (2, Decimal(5), 1.85), (3, Decimal(0), 1.75),
    ]
    assert len(ws._charts) == 3
    chart_rows = []
    for chart in ws._charts:
        assert chart.anchor.ext.cx == cm_to_EMU(TRAJECTORY_CHART_SIZE_CM["width"])
        assert chart.anchor.ext.cy == cm_to_EMU(TRAJECTORY_CHART_SIZE_CM["height"])
        assert chart.anchor._from.col == ord(TRAJECTORY_CHART_ANCHOR_COLUMN) - ord("A")
        assert chart.__class__.__name__ == "ScatterChart"
        assert chart.series[0].graphicalProperties.line.solidFill.srgbClr == "4472C4"
        assert chart.series[0].graphicalProperties.line.width == TRAJECTORY_CHART_LINE_WIDTH_EMU
        assert "1X2" in chart.title.tx.rich.p[0].r[0].t
        assert chart.x_axis.numFmt.formatCode == "0.###"
        assert chart.x_axis.scaling.orientation == "maxMin"
        assert chart.x_axis.scaling.min == 0
        assert chart.x_axis.scaling.max == 120
        assert chart.y_axis.scaling.min == 1
        assert chart.y_axis.scaling.max == 4.5
        chart_rows.append(chart.anchor._from.row)
    assert all(
        next_row - row >= TRAJECTORY_CHART_RESERVED_ROWS
        for row, next_row in zip(chart_rows, chart_rows[1:])
    )
    workbook.close()


def test_charts_share_odds_scale_but_use_choice_specific_time_ranges(tmp_path):
    high_odds = choice("1")
    high_odds.snapshots[1] = replace(
        high_odds.snapshots[1], odds_value=Decimal("8.74"), minutes_before_start=Decimal(240),
    )
    high_odds.snapshots[-1] = replace(
        high_odds.snapshots[-1], minutes_before_start=Decimal("-0.081"),
    )
    low_odds = choice("1")
    workbook = export(tmp_path, context([
        ("1X2", "Full Time", None, "SofaScore", None, [high_odds, choice("X")]),
        ("Asian Handicap", "Full Time", "0", "Pinnacle Sports", None, [low_odds]),
    ]))

    charts = [chart for sheet in workbook for chart in sheet._charts]
    assert len(charts) == 3
    time_ranges = {
        (chart.x_axis.scaling.min, chart.x_axis.scaling.max)
        for chart in charts
    }
    assert time_ranges == {(-0.081, 240), (0, 120)}
    assert all(chart.y_axis.scaling.min == 1 for chart in charts)
    title_to_y_max = {
        chart.title.tx.rich.p[0].r[0].t: chart.y_axis.scaling.max
        for chart in charts
    }
    assert title_to_y_max["1X2 · Full Time · Línea sin línea · Choice 1 · Cuota"] == 9
    assert title_to_y_max["1X2 · Full Time · Línea sin línea · Choice X · Cuota"] == 9
    assert title_to_y_max["Asian Handicap · Full Time · Línea 0 · Choice 1 · Cuota"] == 2
    titles = set(title_to_y_max)
    assert any("1X2" in title for title in titles)
    assert any("Asian Handicap" in title for title in titles)
    workbook.close()


def test_odds_scale_is_shared_by_market_type_period_and_line(tmp_path):
    full_time_main_line = choice("1")
    full_time_main_line.snapshots[1] = replace(
        full_time_main_line.snapshots[1], odds_value=Decimal("8.74"),
    )
    first_half = choice("1")
    first_half.snapshots[0] = replace(
        first_half.snapshots[0], odds_value=Decimal("3.24"),
    )
    workbook = export(tmp_path, context([
        ("1X2", "Full Time", None, "SofaScore", None,
         [full_time_main_line, choice("X")]),
        ("1X2", "Full Time", "0.5", "Pinnacle Sports", None, [choice("1")]),
        ("1X2", "1st Half", None, "SofaScore", None, [first_half]),
    ]))

    worksheet = workbook["Trayectoria 1X2"]
    title_to_y_max = {
        chart.title.tx.rich.p[0].r[0].t: chart.y_axis.scaling.max
        for chart in worksheet._charts
    }
    assert title_to_y_max["1X2 · Full Time · Línea sin línea · Choice 1 · Cuota"] == 9
    assert title_to_y_max["1X2 · Full Time · Línea sin línea · Choice X · Cuota"] == 9
    assert title_to_y_max["1X2 · Full Time · Línea 0.5 · Choice 1 · Cuota"] == 2
    assert title_to_y_max["1X2 · 1st Half · Línea sin línea · Choice 1 · Cuota"] == 3.5
    workbook.close()


def test_locked_workbook_target_saves_to_alternate_path_without_overwriting(tmp_path, monkeypatch):
    output_path = tmp_path / "trajectory.xlsx"
    output_path.write_bytes(b"existing locked workbook")
    real_replace = os.replace

    def replace_with_locked_target(source, destination):
        if Path(destination) == output_path:
            raise PermissionError("simulated Windows file lock")
        return real_replace(source, destination)

    monkeypatch.setattr(excel_exporter.os, "replace", replace_with_locked_target)
    saved_path = excel_exporter.export_odds_trajectory_context_xlsx(
        context([("1X2", "Full Time", None, "SofaScore", None, [choice("1")])]),
        output_path,
    )

    assert saved_path != output_path
    assert "_updated_" in saved_path.name
    assert output_path.read_bytes() == b"existing locked workbook"
    workbook = load_workbook(saved_path)
    assert len(workbook["Trayectoria 1X2"]._charts) == 1
    workbook.close()
